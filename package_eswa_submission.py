#!/usr/bin/env python3
"""Check ESWA limits and copy a flat Editorial Manager folder into submission/."""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SUB = ROOT / "submission"

HIGHLIGHTS = ROOT / "eswa_highlights.txt"
ANON_TEX = ROOT / "eswa_manuscript_anonymous.tex"
IDENT_TEX = ROOT / "eswa_manuscript.tex"

FIGURES = [
    "fig1_application_context.png",
    "fig1_overall_comparison.png",
    "fig2_ablation.png",
    "fig3_horizon_improvement.png",
    "fig4_segment_performance.png",
    "fig6_direction_and_r2.png",
]
EXCLUDED_FIGURE = "fig5_per_window_stability.png"

TEXT_FILES = [
    "eswa_manuscript_anonymous.tex",
    "eswa_title_page.tex",
    "eswa_title_page.txt",
    "eswa_highlights.txt",
    "eswa_cover_letter.txt",
    "eswa_declaration_of_interest.txt",
    "eswa_declaration_of_interest.docx",
    "eswa_data_availability.txt",
    "eswa_credit_statement.txt",
]

# Strings that would identify the author if they appeared in the review file.
FORBIDDEN = [
    "zhanpeng",
    "qin",
    "bytedance",
    "richardandelu",
    "gmail",
    "orcid",
    "0009-0008-9819-8831",
    "qzpthuhhu",
    "github.com",
    "acknowledgment",
    "acknowledgement",
    "competing interest",
    "credit authorship",
    "this research did not receive",
]


def abstract_words(tex: str) -> tuple[int, str]:
    match = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, re.S)
    if not match:
        raise SystemExit("abstract environment not found")
    body = match.group(1)
    body = re.sub(r"\\textbf\{([^}]*)\}", r"\1", body)
    body = re.sub(r"\\emph\{([^}]*)\}", r"\1", body)
    body = body.replace("---", " ")
    body = body.replace("--", "-")
    body = body.replace("\\%", "%")
    body = body.replace("\\&", "&")
    body = re.sub(r"\$([^$]*)\$", r"\1", body)
    body = re.sub(r"\\[a-zA-Z]+", " ", body)
    body = body.replace("{", "").replace("}", "")
    body = re.sub(r"\s+", " ", body).strip()
    return len(body.split()), body


def highlight_lines(path: Path) -> list[str]:
    lines = []
    for raw in path.read_text().splitlines():
        line = raw.strip()
        if not line or line.lower().startswith("highlights"):
            continue
        if line.startswith("- "):
            line = line[2:].strip()
        lines.append(line)
    return lines


def check() -> dict:
    anon = ANON_TEX.read_text()
    ident = IDENT_TEX.read_text()
    anon_n, anon_abs = abstract_words(anon)
    ident_n, ident_abs = abstract_words(ident)
    if anon_abs != ident_abs:
        raise SystemExit("abstract text differs between identified and anonymized manuscripts")
    if anon_n > 250:
        raise SystemExit(f"abstract has {anon_n} words (limit 250)")
    highlights = highlight_lines(HIGHLIGHTS)
    if not 3 <= len(highlights) <= 5:
        raise SystemExit(f"expected 3-5 highlights, found {len(highlights)}")
    for line in highlights:
        if len(line) > 85:
            raise SystemExit(f"highlight has {len(line)} characters (limit 85): {line}")
    low = anon.lower()
    for token in FORBIDDEN:
        if token in low:
            raise SystemExit(f"anonymized manuscript contains identifying text: {token}")
    if "\\includegraphics" not in anon:
        raise SystemExit("anonymized manuscript has no figures")
    for match in re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", anon):
        if "/" in match or "\\" in match:
            raise SystemExit(f"figure path is not flat: {match}")
    if EXCLUDED_FIGURE in anon or "fig5" in anon:
        raise SystemExit("fig5 is referenced; it is simulated and must stay out of the manuscript")
    letter = (ROOT / "eswa_cover_letter.txt").read_text()
    if re.search(r"\bWe\b|\bOur\b|\bAll authors\b", letter):
        raise SystemExit("cover letter still uses plural author voice")
    funding = "This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors."
    title = (ROOT / "eswa_title_page.txt").read_text()
    if funding not in title or funding not in ident:
        raise SystemExit("official funding sentence missing from title page or identified manuscript")
    if funding in anon:
        raise SystemExit("funding sentence leaked into the anonymized manuscript")
    return {
        "abstract_words": anon_n,
        "highlights": [(len(line), line) for line in highlights],
    }


def copy_flat() -> None:
    if SUB.exists():
        shutil.rmtree(SUB)
    SUB.mkdir()
    for name in TEXT_FILES:
        src = ROOT / name
        if not src.exists():
            raise SystemExit(f"missing {name}")
        shutil.copy2(src, SUB / name)
    for name in FIGURES:
        src = ROOT / "figures" / name
        if not src.exists():
            raise SystemExit(f"missing figure {name}")
        shutil.copy2(src, SUB / name)
    excluded = SUB / EXCLUDED_FIGURE
    if excluded.exists():
        raise SystemExit("fig5 was copied into submission/")


def write_declaration_docx() -> None:
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    title = doc.add_paragraph()
    run = title.add_run("Declaration of Interest Statement")
    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph("Expert Systems with Applications")
    doc.add_paragraph(
        "Manuscript title: When Feature Engineering Outperforms Deep Learning: "
        "Sparsity-Aware Direct Multi-Step Forecasting for Small-Scale Retail Demand"
    )
    doc.add_paragraph("Author: Zhanpeng Qin")
    doc.add_paragraph("Date: 14 October 2026")

    heading = doc.add_paragraph()
    heading.add_run("Declaration of interests").bold = True
    doc.add_paragraph(
        "\u2612 The author declares that they have no known competing financial interests "
        "or personal relationships that could have appeared to influence the work reported "
        "in this paper."
    )
    doc.add_paragraph(
        "\u2610 The author declares the following financial interests/personal relationships "
        "which may be considered as potential competing interests:"
    )
    doc.add_paragraph("Signed: Zhanpeng Qin")
    doc.add_paragraph("ByteDance, Beijing, China")
    doc.add_paragraph("richardandelu50@gmail.com")
    doc.add_paragraph("ORCID: 0009-0008-9819-8831")
    doc.save(ROOT / "eswa_declaration_of_interest.docx")


def compile_tex(workdir: Path, tex_name: str, output_dir: Path | None = None) -> Path:
    out = output_dir or workdir
    out.mkdir(parents=True, exist_ok=True)
    stem = Path(tex_name).stem
    cmd = ["pdflatex", "-interaction=nonstopmode"]
    if output_dir is not None:
        cmd.append(f"-output-directory={output_dir}")
    cmd.append(tex_name)
    for _ in range(2):
        proc = subprocess.run(cmd, cwd=workdir, capture_output=True, text=True)
        pdf = out / f"{stem}.pdf"
        if proc.returncode != 0 or not pdf.exists():
            tail = (proc.stdout or "")[-4000:]
            raise SystemExit(f"pdflatex failed for {tex_name} in {workdir}\n{tail}")
    log_path = out / f"{stem}.log"
    if log_path.exists():
        log = log_path.read_text(errors="replace")
        fatal = [ln for ln in log.splitlines() if ln.startswith("!")]
        if fatal:
            raise SystemExit(f"LaTeX errors in {tex_name}:\n" + "\n".join(fatal[:30]))
        undefined = [
            ln for ln in log.splitlines()
            if "undefined" in ln.lower() and ("reference" in ln.lower() or "citation" in ln.lower())
        ]
        if undefined:
            raise SystemExit(f"Unresolved references in {tex_name}:\n" + "\n".join(undefined[:20]))
    for suffix in (".aux", ".log", ".out", ".bbl", ".blg", ".spl"):
        leftover = out / f"{stem}{suffix}"
        if leftover.exists():
            leftover.unlink()
    return out / f"{stem}.pdf"


def assert_pdf_text_clean(pdf: Path) -> None:
    extracted = pdf.with_suffix(".extracted.txt")
    proc = subprocess.run(["pdftotext", str(pdf), str(extracted)], capture_output=True, text=True)
    if proc.returncode != 0:
        raise SystemExit(f"pdftotext failed for {pdf}: {proc.stderr}")
    text = extracted.read_text(errors="replace").lower()
    extracted.unlink()
    for token in FORBIDDEN:
        if token in text:
            raise SystemExit(f"anonymized PDF contains identifying text: {token}")


def main() -> None:
    write_declaration_docx()
    info = check()
    copy_flat()
    anon_pdf = compile_tex(SUB, "eswa_manuscript_anonymous.tex")
    compile_tex(SUB, "eswa_title_page.tex")
    assert_pdf_text_clean(anon_pdf)
    # Identified author copy is compiled only to prove it builds. The PDF is not
    # placed in submission/, because that folder is the review upload.
    build = ROOT / ".latex-build"
    if build.exists():
        shutil.rmtree(build)
    compile_tex(ROOT, IDENT_TEX.name, output_dir=build)
    shutil.rmtree(build)
    print(f"abstract_words={info['abstract_words']}")
    for n, line in info["highlights"]:
        print(f"highlight_{n}={line}")
    print(f"submission={SUB}")


if __name__ == "__main__":
    main()
